"""
Shared Document Text Layer

One place that turns a file on disk into text. Every scanner, classifier, and
extractor in the dataroom pipeline reads through here instead of importing
pypdf/pdfplumber/openpyxl directly.

Three properties the per-extractor implementations did not have:

1. **Loud failure.** A missing parser library or an unreadable file produces an
   ``ExtractedText`` whose ``error`` is set and whose ``ok`` is False. Callers
   can surface it. The old code wrapped every parse in a bare ``except`` and
   returned ``""``, so a dataroom with an uninstalled ``pypdf`` looked exactly
   like a dataroom full of empty PDFs.

2. **One cache.** A 100MB deck used by the team, traction, and competitive
   extractors is parsed once per process, not three times.

3. **Scan detection.** Executed financing documents are frequently scans. A PDF
   that yields almost no characters per page is flagged ``is_scanned`` so the
   caller knows the silence is a scanner artifact rather than an empty file.
"""

from __future__ import annotations

import email
import email.policy
import re
import subprocess
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional


# =============================================================================
# Result type
# =============================================================================

@dataclass
class ExtractedText:
    """Text pulled from one file, plus how it was obtained and what went wrong."""

    path: str
    text: str = ""
    method: str = "none"           # which backend produced the text
    page_count: Optional[int] = None
    char_count: int = 0
    truncated: bool = False
    is_scanned: bool = False       # PDF with pages but almost no extractable text
    error: Optional[str] = None
    notes: List[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.error is None and bool(self.text.strip())

    def head(self, n: int = 4000) -> str:
        return self.text[:n]


# Below this many characters per page, a PDF is almost certainly a scan.
SCANNED_CHARS_PER_PAGE = 60

# Guard against a single pathological file eating the whole context budget.
DEFAULT_MAX_CHARS = 400_000

_CACHE: Dict[tuple, ExtractedText] = {}


# =============================================================================
# Public API
# =============================================================================

def extract_text(
    path: str | Path,
    max_chars: int = DEFAULT_MAX_CHARS,
    use_cache: bool = True,
    ocr: bool = True,
    ocr_max_pages: int = 25,
) -> ExtractedText:
    """
    Extract text from any supported document.

    Never raises for a parse failure — inspect ``.ok`` and ``.error`` instead.

    Args:
        ocr: Run OCR when a PDF turns out to be a scan. Executed financing
            documents are routinely scanned, so this defaults on; pass False
            for a fast inventory pass where only the readable files matter.
        ocr_max_pages: Cap on pages sent through Tesseract per document.
    """
    p = Path(path)

    if not p.exists():
        return ExtractedText(path=str(p), error="file does not exist")

    try:
        stat = p.stat()
        key = (str(p.resolve()), stat.st_mtime_ns, stat.st_size, max_chars, ocr)
    except OSError as e:
        return ExtractedText(path=str(p), error=f"cannot stat file: {e}")

    if use_cache and key in _CACHE:
        return _CACHE[key]

    handler = _HANDLERS.get(p.suffix.lower())
    if handler is None:
        result = ExtractedText(
            path=str(p),
            error=f"no text handler for extension '{p.suffix.lower()}'",
        )
    else:
        try:
            result = handler(p)
        except Exception as e:  # a backend blew up in a way we did not anticipate
            result = ExtractedText(
                path=str(p),
                error=f"{type(e).__name__}: {e}",
            )

    result = _finalize(result, max_chars)

    # A scan carries its content in pixels. Recover it before the caller sees
    # an empty-looking document and concludes the file had nothing in it.
    if ocr and p.suffix.lower() == ".pdf" and result.is_scanned:
        result = _ocr_pdf(p, result, max_pages=ocr_max_pages)
        result = _finalize(result, max_chars)

    if use_cache:
        _CACHE[key] = result
    return result


def extract_many(paths, max_chars: int = DEFAULT_MAX_CHARS, **kw) -> List[ExtractedText]:
    """Extract text from several files, preserving order."""
    return [extract_text(p, max_chars=max_chars, **kw) for p in paths]


def supported_extensions() -> set:
    """Extensions this module can turn into text."""
    return set(_HANDLERS)


def clear_cache() -> None:
    _CACHE.clear()


def _finalize(result: ExtractedText, max_chars: int) -> ExtractedText:
    """Normalize whitespace, apply the cap, and flag probable scans."""
    if result.text:
        # Collapse the ragged runs of blank lines that PDF backends emit.
        result.text = re.sub(r"\n{4,}", "\n\n\n", result.text).strip()

    if max_chars and len(result.text) > max_chars:
        result.text = result.text[:max_chars]
        result.truncated = True
        result.notes.append(f"truncated to {max_chars} characters")

    result.char_count = len(result.text)

    if result.page_count and result.page_count > 0:
        per_page = result.char_count / result.page_count
        if per_page < SCANNED_CHARS_PER_PAGE:
            result.is_scanned = True
            result.notes.append(
                f"only {per_page:.0f} chars/page — likely a scanned document; "
                f"OCR required for full text"
            )

    return result


# =============================================================================
# PDF
# =============================================================================

def _extract_pdf(p: Path) -> ExtractedText:
    """
    PDF text via PyMuPDF, falling back to pdfplumber then pypdf.

    PyMuPDF leads because it is both the fastest and the most tolerant of the
    malformed PDFs that e-signature platforms produce.
    """
    attempts: List[str] = []

    # --- PyMuPDF ---
    try:
        import pymupdf

        with pymupdf.open(str(p)) as doc:
            unlocked_with = None
            if doc.needs_pass:
                unlocked_with = _try_passwords(doc, p)
                if unlocked_with is None:
                    return ExtractedText(
                        path=str(p),
                        page_count=doc.page_count,
                        error="PDF is password protected and no password was recoverable",
                    )
            pages = [page.get_text("text") for page in doc]
            result = ExtractedText(
                path=str(p),
                text="\n\n".join(pages),
                method="pymupdf",
                page_count=len(pages),
            )
            if unlocked_with:
                result.notes.append(f"unlocked with password from filename: {unlocked_with!r}")
            return result
    except ImportError:
        attempts.append("pymupdf not installed")
    except Exception as e:
        attempts.append(f"pymupdf: {type(e).__name__}: {e}")

    # --- pdfplumber ---
    try:
        import pdfplumber

        with pdfplumber.open(str(p)) as pdf:
            pages = [(page.extract_text() or "") for page in pdf.pages]
            return ExtractedText(
                path=str(p),
                text="\n\n".join(pages),
                method="pdfplumber",
                page_count=len(pages),
                notes=attempts,
            )
    except ImportError:
        attempts.append("pdfplumber not installed")
    except Exception as e:
        attempts.append(f"pdfplumber: {type(e).__name__}: {e}")

    # --- pypdf ---
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(p))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return ExtractedText(
            path=str(p),
            text="\n\n".join(pages),
            method="pypdf",
            page_count=len(pages),
            notes=attempts,
        )
    except ImportError:
        attempts.append("pypdf not installed")
    except Exception as e:
        attempts.append(f"pypdf: {type(e).__name__}: {e}")

    return ExtractedText(
        path=str(p),
        error="all PDF backends failed: " + "; ".join(attempts),
    )


def _try_passwords(doc, p: Path) -> Optional[str]:
    """
    Attempt to open a locked PDF with passwords implied by its own filename.

    Sharing a board deck as ``<name> - pwd SECRET.pdf`` is a real and common
    habit; the password travels with the file precisely so the recipient can
    open it. Reading it off the filename is not an attack on the protection, it
    is using the credential the sender attached. Anything not derivable from
    the filename stays locked and is reported as such.
    """
    name = p.stem
    candidates: List[str] = []

    # "... - pwd SECRET", "password: SECRET", "pw=SECRET"
    for m in re.finditer(
        r"(?:pwd|pw|password|passcode)\s*[:=\-]?\s*([A-Za-z0-9._!@#-]{4,})",
        name,
        re.IGNORECASE,
    ):
        candidates.append(m.group(1))

    # The trailing token of a name that mentions a password at all.
    if re.search(r"(?:pwd|pw|password|passcode)", name, re.IGNORECASE):
        tail = name.split()[-1]
        if len(tail) >= 4:
            candidates.append(tail)

    for candidate in dict.fromkeys(candidates):  # de-dup, keep order
        try:
            if doc.authenticate(candidate):
                return candidate
        except Exception:
            continue
    return None


# =============================================================================
# OCR
# =============================================================================

# Tesseract and MuPDF both refuse to allocate past a point. A page rendered
# above this many pixels is downscaled rather than attempted.
OCR_MAX_PIXELS = 40_000_000

# Some PDFs are whiteboard or canvas exports hundreds of inches across. They
# cannot be rendered at a resolution where OCR would mean anything, and
# attempting it costs minutes per page.
OVERSIZED_PAGE_POINTS = 20_000


def _ocr_pdf(p: Path, base: ExtractedText, max_pages: int = 25) -> ExtractedText:
    """
    Recover text from a scanned PDF by rendering pages and running Tesseract.

    Preserves whatever the text layer did yield; a scan with a partial text
    layer (common when a signature page is appended to a born-digital document)
    should end up with both.
    """
    try:
        import pymupdf
    except ImportError:
        base.notes.append("OCR skipped: pymupdf not installed")
        return base

    if not _have_tesseract():
        base.notes.append("OCR skipped: tesseract binary not found on PATH")
        return base

    pieces: List[str] = []
    failures = 0

    try:
        doc = pymupdf.open(str(p))
    except Exception as e:
        base.notes.append(f"OCR skipped: {type(e).__name__}: {e}")
        return base

    with doc:
        if doc.needs_pass and _try_passwords(doc, p) is None:
            base.notes.append("OCR skipped: document is locked")
            return base

        for index, page in enumerate(doc):
            if index >= max_pages:
                base.notes.append(
                    f"OCR stopped after {max_pages} of {doc.page_count} pages"
                )
                break

            rect = page.rect
            if max(rect.width, rect.height) > OVERSIZED_PAGE_POINTS:
                base.notes.append(
                    f"page {index + 1} is an oversized canvas "
                    f"({rect.width:.0f}x{rect.height:.0f}pt) — likely a whiteboard or "
                    f"diagram export; OCR would not produce meaningful text"
                )
                continue

            text = _ocr_page(page, pymupdf)
            if text is None:
                failures += 1
            else:
                pieces.append(text)

    if not pieces:
        if failures:
            base.notes.append(f"OCR produced no text ({failures} pages failed to render)")
        return base

    recovered = "\n\n".join(pieces)
    base.text = (base.text + "\n\n" + recovered).strip() if base.text.strip() else recovered
    base.method = f"{base.method}+tesseract"
    base.notes.append(f"OCR recovered {len(recovered)} characters from {len(pieces)} page(s)")
    if failures:
        base.notes.append(f"{failures} page(s) could not be rendered for OCR")
    # The text is real now even though the source was a scan; keep the flag so
    # downstream callers can weight OCR text as lower-fidelity than a text layer.
    return base


def _ocr_page(page, pymupdf) -> Optional[str]:
    """Render one page at the largest scale that fits the pixel budget, then OCR."""
    import os
    import tempfile

    for scale in (2.0, 1.5, 1.0, 0.6):
        try:
            pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale))
            if pix.width * pix.height > OCR_MAX_PIXELS:
                continue
            tmp_path = None
            try:
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                    tmp.write(pix.tobytes("png"))
                    tmp_path = tmp.name
                out = subprocess.run(
                    ["tesseract", tmp_path, "stdout", "--psm", "3"],
                    capture_output=True,
                    timeout=180,
                )
                if out.returncode == 0:
                    return out.stdout.decode("utf-8", errors="replace")
            finally:
                if tmp_path and os.path.exists(tmp_path):
                    os.unlink(tmp_path)
        except Exception:
            continue
    return None


_TESSERACT_PRESENT: Optional[bool] = None


def _have_tesseract() -> bool:
    global _TESSERACT_PRESENT
    if _TESSERACT_PRESENT is None:
        import shutil

        _TESSERACT_PRESENT = shutil.which("tesseract") is not None
    return _TESSERACT_PRESENT


def extract_pdf_tables(p: str | Path, max_pages: int = 40) -> List[List[List[str]]]:
    """
    Pull ruled tables out of a PDF via pdfplumber.

    Separate from ``extract_text`` because table extraction is an order of
    magnitude slower and only the cap-table and financial extractors want it.
    Returns a list of tables; each table is a list of rows of cell strings.
    """
    try:
        import pdfplumber
    except ImportError:
        return []

    tables: List[List[List[str]]] = []
    try:
        with pdfplumber.open(str(p)) as pdf:
            for page in pdf.pages[:max_pages]:
                for raw in page.extract_tables() or []:
                    cleaned = [
                        [(cell or "").strip() for cell in row]
                        for row in raw
                        if any((cell or "").strip() for cell in row)
                    ]
                    if cleaned:
                        tables.append(cleaned)
    except Exception:
        return tables
    return tables


# =============================================================================
# Word
# =============================================================================

def _extract_docx(p: Path) -> ExtractedText:
    """
    DOCX text including table cells.

    Tables matter disproportionately here: schedules of purchasers, disclosure
    schedules, and pro-forma exhibits live in tables, and a paragraph-only walk
    silently drops all of them.
    """
    try:
        import docx
    except ImportError:
        return ExtractedText(path=str(p), error="python-docx not installed")

    document = docx.Document(str(p))
    chunks: List[str] = [para.text for para in document.paragraphs if para.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))

    return ExtractedText(path=str(p), text="\n".join(chunks), method="python-docx")


def _extract_doc(p: Path) -> ExtractedText:
    """Legacy .doc via macOS textutil — the only converter we can rely on here."""
    try:
        out = subprocess.run(
            ["textutil", "-convert", "txt", "-stdout", str(p)],
            capture_output=True, timeout=60,
        )
        if out.returncode == 0:
            return ExtractedText(
                path=str(p),
                text=out.stdout.decode("utf-8", errors="replace"),
                method="textutil",
            )
        return ExtractedText(
            path=str(p),
            error=f"textutil failed: {out.stderr.decode('utf-8', errors='replace')[:200]}",
        )
    except FileNotFoundError:
        return ExtractedText(path=str(p), error="legacy .doc needs macOS textutil")
    except subprocess.TimeoutExpired:
        return ExtractedText(path=str(p), error="textutil timed out")


# =============================================================================
# Spreadsheets
# =============================================================================

def _extract_xlsx(p: Path) -> ExtractedText:
    """
    Flatten every sheet to pipe-delimited rows under a sheet header.

    Formula cells resolve to their cached value (``data_only=True``); a workbook
    that has never been opened in Excel yields None for those, which we render
    as empty rather than as the formula string.
    """
    try:
        import openpyxl
    except ImportError:
        return ExtractedText(path=str(p), error="openpyxl not installed")

    wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
    chunks: List[str] = []
    try:
        for sheet in wb.worksheets:
            chunks.append(f"\n=== SHEET: {sheet.title} ===")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                if any(cells):
                    chunks.append(" | ".join(cells))
    finally:
        wb.close()

    return ExtractedText(path=str(p), text="\n".join(chunks), method="openpyxl")


def _extract_xls(p: Path) -> ExtractedText:
    """Legacy .xls — openpyxl cannot read it; xlrd can if present."""
    try:
        import xlrd
    except ImportError:
        return ExtractedText(path=str(p), error="legacy .xls needs xlrd")

    book = xlrd.open_workbook(str(p))
    chunks: List[str] = []
    for sheet in book.sheets():
        chunks.append(f"\n=== SHEET: {sheet.name} ===")
        for r in range(sheet.nrows):
            cells = [str(c).strip() for c in sheet.row_values(r)]
            if any(cells):
                chunks.append(" | ".join(cells))
    return ExtractedText(path=str(p), text="\n".join(chunks), method="xlrd")


def _extract_csv(p: Path) -> ExtractedText:
    text = p.read_text(encoding="utf-8", errors="replace")
    return ExtractedText(path=str(p), text=text, method="read_text")


# =============================================================================
# Presentations
# =============================================================================

def _extract_pptx(p: Path) -> ExtractedText:
    """Slide text with a per-slide header, including speaker notes and tables."""
    try:
        from pptx import Presentation
    except ImportError:
        return ExtractedText(path=str(p), error="python-pptx not installed")

    prs = Presentation(str(p))
    chunks: List[str] = []
    slide_count = 0

    for idx, slide in enumerate(prs.slides, start=1):
        slide_count = idx
        chunks.append(f"\n=== SLIDE {idx} ===")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunks.append(shape.text_frame.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        chunks.append(" | ".join(cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                chunks.append(f"[speaker notes] {notes}")

    return ExtractedText(
        path=str(p),
        text="\n".join(chunks),
        method="python-pptx",
        page_count=slide_count,
    )


# =============================================================================
# Mail and Apple bundles
# =============================================================================

def _extract_eml(p: Path) -> ExtractedText:
    """
    RFC-822 message: headers that establish provenance, then the body.

    Closing-confirmation emails are frequently the only artifact that dates a
    financing, so the headers are part of the payload, not chrome.
    """
    raw = p.read_bytes()
    msg = email.message_from_bytes(raw, policy=email.policy.default)

    header_lines = [
        f"{h}: {msg.get(h)}"
        for h in ("Date", "From", "To", "Cc", "Subject")
        if msg.get(h)
    ]

    body = ""
    try:
        part = msg.get_body(preferencelist=("plain", "html"))
        if part is not None:
            body = part.get_content()
            if part.get_content_type() == "text/html":
                body = re.sub(r"<[^>]+>", " ", body)
                body = re.sub(r"[ \t]{2,}", " ", body)
    except Exception as e:
        body = f"[body could not be decoded: {e}]"

    attachments = [
        a.get_filename()
        for a in msg.iter_attachments()
        if a.get_filename()
    ]
    if attachments:
        header_lines.append("Attachments: " + ", ".join(attachments))

    return ExtractedText(
        path=str(p),
        text="\n".join(header_lines) + "\n\n" + body,
        method="email.parser",
    )


def _extract_pages(p: Path) -> ExtractedText:
    """
    Apple Pages bundle.

    Older bundles carry a rendered ``QuickLook/Preview.pdf`` we can read
    directly. Newer ones store only Snappy-compressed protobuf (``.iwa``), which
    is not worth reimplementing — those we report honestly so the operator knows
    to export a PDF rather than assuming the file was empty.
    """
    try:
        with zipfile.ZipFile(str(p)) as zf:
            names = zf.namelist()
            preview = next(
                (n for n in names if n.lower().endswith("quicklook/preview.pdf")), None
            )
            if preview:
                import tempfile

                with tempfile.NamedTemporaryFile(suffix=".pdf", delete=True) as tmp:
                    tmp.write(zf.read(preview))
                    tmp.flush()
                    inner = _extract_pdf(Path(tmp.name))
                    inner.path = str(p)
                    inner.method = f"pages-quicklook+{inner.method}"
                    return inner

            if any(n.endswith(".iwa") for n in names):
                return ExtractedText(
                    path=str(p),
                    error=(
                        "Apple Pages bundle in IWA format with no QuickLook preview — "
                        "export to PDF or DOCX to make this readable"
                    ),
                )
    except zipfile.BadZipFile:
        return ExtractedText(path=str(p), error="not a readable Pages bundle")

    return ExtractedText(path=str(p), error="unrecognized Pages bundle layout")


# =============================================================================
# Plain text
# =============================================================================

def _extract_plain(p: Path) -> ExtractedText:
    return ExtractedText(
        path=str(p),
        text=p.read_text(encoding="utf-8", errors="replace"),
        method="read_text",
    )


_HANDLERS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".doc": _extract_doc,
    ".xlsx": _extract_xlsx,
    ".xlsm": _extract_xlsx,
    ".xls": _extract_xls,
    ".csv": _extract_csv,
    ".pptx": _extract_pptx,
    ".eml": _extract_eml,
    ".pages": _extract_pages,
    ".md": _extract_plain,
    ".txt": _extract_plain,
    ".rtf": _extract_doc,
}
