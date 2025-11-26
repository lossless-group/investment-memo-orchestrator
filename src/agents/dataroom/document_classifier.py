"""
Document Classifier for Dataroom Analysis

Three-stage classification approach:
1. Directory-based pre-classification (highest confidence)
2. Filename heuristic classification
3. LLM-based content classification (only if needed)
"""

from pathlib import Path
from typing import List, Tuple, Optional
import re
import json

from .dataroom_state import DocumentInventoryItem
from .document_scanner import parse_directory_category


# =============================================================================
# Directory Category Mapping
# =============================================================================

DIRECTORY_CATEGORY_MAP = {
    # Executive / Pitch
    "executive summary": "pitch_deck",
    "executive": "pitch_deck",
    "pitch": "pitch_deck",
    "deck": "pitch_deck",
    "investor": "pitch_deck",

    # Product
    "product overview": "product_documentation",
    "product": "product_documentation",
    "technology": "product_documentation",
    "architecture": "product_documentation",

    # Marketing
    "marketing materials": "marketing_collateral",
    "marketing": "marketing_collateral",
    "collateral": "marketing_collateral",
    "sales materials": "marketing_collateral",

    # Competitive
    "gtm": "competitive_analysis",
    "gtm competitive": "competitive_analysis",
    "competitive overview": "competitive_analysis",
    "competitive": "competitive_analysis",
    "compete": "competitive_analysis",
    "battlecard": "competitive_analysis",
    "battlecards": "competitive_analysis",

    # Financial
    "financial overview": "financial_statements",
    "financial": "financial_statements",
    "financials": "financial_statements",
    "finance": "financial_statements",

    # Legal
    "legal": "legal_document",
    "legal documents": "legal_document",
    "corporate": "legal_document",

    # Team
    "team": "team_bios",
    "leadership": "team_bios",
    "management": "team_bios",
    "founders": "team_bios",

    # Traction
    "traction": "traction_metrics",
    "customers": "customer_list",
    "pipeline": "pipeline_metrics",
    "metrics": "traction_metrics",

    # Governance
    "governance": "investor_update",
    "board": "investor_update",
    "updates": "investor_update",
}


# =============================================================================
# Filename Pattern Mapping
# =============================================================================

FILENAME_PATTERNS = {
    "pitch_deck": [
        r"pitch.*deck", r"investor.*deck", r"presentation",
        r"one.*pager", r"teaser", r"executive.*summary",
        r"deck.*\d{4}", r"expanded.*deck",
    ],
    "competitive_analysis": [
        r"competitive", r"competitor", r"landscape", r"battlecard",
        r"battle.*card", r"comparison", r"vs\b", r"versus",
        r"market.*map", r"swot", r"compete",
    ],
    "financial_statements": [
        r"p&l", r"p\s*&\s*l", r"profit.*loss", r"income.*statement",
        r"balance.*sheet", r"cash.*flow", r"financials?$",
    ],
    "financial_projections": [
        r"model", r"projection", r"forecast", r"operating.*model",
        r"\d{4}.*model", r"financial.*model", r"budget",
    ],
    "cap_table": [
        r"cap.*table", r"captable", r"ownership", r"equity",
        r"shareholding", r"stock.*ledger",
    ],
    "term_sheet": [
        r"term.*sheet", r"termsheet", r"terms",
        r"series.*[a-z].*terms",
    ],
    "safe_note": [
        r"\bsafe\b", r"simple.*agreement", r"convertible",
        r"note.*purchase",
    ],
    "team_bios": [
        r"\bteam\b", r"\bbio", r"founder", r"leadership",
        r"management", r"executive.*team",
    ],
    "customer_list": [
        r"customer", r"client", r"account", r"logo",
    ],
    "product_documentation": [
        r"product", r"spec", r"architecture", r"roadmap",
        r"diagram", r"technical",
    ],
    "marketing_collateral": [
        r"marketing", r"collateral", r"datasheet", r"brochure",
        r"one.*pager", r"integration.*guide",
    ],
}


# =============================================================================
# Content Signal Patterns
# =============================================================================

CONTENT_SIGNALS = {
    "pitch_deck": [
        "investment opportunity", "use of funds", "ask",
        "tam", "sam", "som", "market size",
        "traction", "milestones", "roadmap",
    ],
    "competitive_analysis": [
        "competitors", "competitive landscape", "market share",
        "differentiation", "positioning", "strengths", "weaknesses",
        "swot", "feature comparison", "pricing comparison",
        "competitive advantage", "barriers to entry", "threat",
        "winning angle", "discovery questions",
    ],
    "financial_statements": [
        "revenue", "cogs", "gross margin", "gross profit",
        "operating expenses", "ebitda", "net income",
        "assets", "liabilities", "equity",
    ],
    "cap_table": [
        "shares outstanding", "fully diluted", "ownership %",
        "option pool", "preferred stock", "common stock",
        "vesting", "strike price",
    ],
    "term_sheet": [
        "pre-money valuation", "post-money", "price per share",
        "liquidation preference", "anti-dilution",
        "board composition", "protective provisions",
    ],
}


# =============================================================================
# Main Classification Functions
# =============================================================================

def classify_documents(
    inventory: List[DocumentInventoryItem],
    use_llm: bool = False  # Default to False to avoid API calls during testing
) -> List[DocumentInventoryItem]:
    """
    Classify all documents in inventory using three-stage approach.

    Args:
        inventory: List of DocumentInventoryItem objects
        use_llm: Whether to use LLM for uncertain classifications

    Returns:
        Updated inventory with classifications
    """
    for item in inventory:
        # Stage 1: Directory-based classification
        doc_type, confidence, source = _classify_by_directory(item)

        if confidence >= 0.8:
            item["document_type"] = doc_type
            item["classification_confidence"] = confidence
            item["classification_source"] = source
            item["classification_reasoning"] = f"Directory match: {item['parent_directory']}"

            # Refine with filename patterns (e.g., cap table in financial folder)
            item = refine_classification(item)
            continue

        # Stage 2: Filename-based classification
        doc_type_fn, confidence_fn = _classify_by_filename(item)

        if confidence_fn > confidence:
            doc_type = doc_type_fn
            confidence = confidence_fn
            source = "filename"

        if confidence >= 0.7:
            item["document_type"] = doc_type
            item["classification_confidence"] = confidence
            item["classification_source"] = source
            item["classification_reasoning"] = f"Filename match: {item['filename']}"
            continue

        # Stage 3: LLM-based classification (optional)
        if use_llm and confidence < 0.7:
            doc_type_llm, confidence_llm, reasoning = _classify_by_content(item)
            if confidence_llm > confidence:
                doc_type = doc_type_llm
                confidence = confidence_llm
                source = "content"
                item["classification_reasoning"] = reasoning

        # Apply best classification found
        item["document_type"] = doc_type
        item["classification_confidence"] = confidence
        item["classification_source"] = source

        if not item["classification_reasoning"]:
            item["classification_reasoning"] = "Low confidence classification"

    return inventory


def _classify_by_directory(item: DocumentInventoryItem) -> Tuple[str, float, str]:
    """
    Stage 1: Classify based on parent directory name.

    Returns:
        (document_type, confidence, source)
    """
    parent_dir = item["parent_directory"]
    category = parse_directory_category(parent_dir)

    if not category:
        return "unknown", 0.0, "unknown"

    # Check for exact match
    if category in DIRECTORY_CATEGORY_MAP:
        return DIRECTORY_CATEGORY_MAP[category], 0.9, "directory"

    # Check for partial match
    for key, doc_type in DIRECTORY_CATEGORY_MAP.items():
        if key in category or category in key:
            return doc_type, 0.85, "directory"

    return "unknown", 0.0, "unknown"


def _classify_by_filename(item: DocumentInventoryItem) -> Tuple[str, float]:
    """
    Stage 2: Classify based on filename patterns.

    Returns:
        (document_type, confidence)
    """
    filename_lower = item["filename"].lower()

    best_type = "unknown"
    best_confidence = 0.0

    for doc_type, patterns in FILENAME_PATTERNS.items():
        for pattern in patterns:
            if re.search(pattern, filename_lower):
                # Give higher confidence to longer/more specific matches
                match = re.search(pattern, filename_lower)
                if match:
                    match_length = len(match.group())
                    confidence = min(0.6 + (match_length / 30), 0.85)
                    if confidence > best_confidence:
                        best_type = doc_type
                        best_confidence = confidence

    # Extension-based hints (lower confidence)
    ext = item["extension"]
    if best_confidence < 0.5:
        if ext == ".csv":
            # CSV files are often financial models
            best_type = "financial_projections"
            best_confidence = 0.5
        elif ext in [".xlsx", ".xls"]:
            # Excel could be financials or cap table
            if "cap" in filename_lower or "table" in filename_lower:
                best_type = "cap_table"
                best_confidence = 0.6
            else:
                best_type = "financial_statements"
                best_confidence = 0.4

    return best_type, best_confidence


def refine_classification(item: DocumentInventoryItem) -> DocumentInventoryItem:
    """
    Refine classification using filename patterns after directory-based classification.

    This catches cases where directory gives broad category but filename is more specific
    (e.g., "Financial Overview" directory but "Cap Table" filename).
    """
    filename_lower = item["filename"].lower()

    # Cap table detection (overrides financial_statements)
    if item["document_type"] == "financial_statements":
        if "cap" in filename_lower and "table" in filename_lower:
            item["document_type"] = "cap_table"
            item["classification_reasoning"] = "Filename indicates cap table"
        elif "operating" in filename_lower and "model" in filename_lower:
            item["document_type"] = "financial_projections"
            item["classification_reasoning"] = "Filename indicates financial projections/model"

    return item


def _classify_by_content(item: DocumentInventoryItem) -> Tuple[str, float, str]:
    """
    Stage 3: Classify based on document content using LLM.

    Returns:
        (document_type, confidence, reasoning)
    """
    # Extract content sample
    content_sample = _extract_content_sample(item["file_path"])

    if not content_sample:
        return "unknown", 0.3, "Could not extract content for classification"

    try:
        from langchain_anthropic import ChatAnthropic

        llm = ChatAnthropic(
            model="claude-sonnet-4-5-20250929",
            temperature=0,
            max_tokens=500
        )

        # Build category list
        categories = list(set(DIRECTORY_CATEGORY_MAP.values()))
        category_list = "\n".join(f"- {cat}" for cat in sorted(categories))

        prompt = f"""Classify this document into ONE of these categories:
{category_list}

FILENAME: {item["filename"]}
EXTENSION: {item["extension"]}

CONTENT SAMPLE (first ~2000 chars):
---
{content_sample[:2000]}
---

Respond with JSON only:
{{"document_type": "<category>", "confidence": <0.0-1.0>, "reasoning": "<brief explanation>"}}
"""

        response = llm.invoke(prompt)
        result = json.loads(response.content)

        return (
            result.get("document_type", "unknown"),
            result.get("confidence", 0.5),
            result.get("reasoning", "LLM classification")
        )

    except Exception as e:
        return "unknown", 0.3, f"LLM classification failed: {str(e)}"


def _extract_content_sample(file_path: str) -> Optional[str]:
    """Extract text sample from document for classification."""
    path = Path(file_path)
    extension = path.suffix.lower()

    try:
        if extension == ".pdf":
            return _extract_pdf_sample(path)
        elif extension in [".xlsx", ".xls"]:
            return _extract_excel_sample(path)
        elif extension == ".csv":
            return _extract_csv_sample(path)
        elif extension == ".docx":
            return _extract_docx_sample(path)
        elif extension in [".txt", ".md"]:
            return _extract_text_sample(path)
        else:
            return None
    except Exception as e:
        return f"[Extraction error: {str(e)}]"


def _extract_pdf_sample(path: Path) -> Optional[str]:
    """Extract text from first 3 pages of PDF."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        text_parts = []
        for i, page in enumerate(reader.pages[:3]):
            text = page.extract_text()
            if text:
                text_parts.append(f"--- PAGE {i+1} ---\n{text}")
        return "\n".join(text_parts)
    except Exception:
        return None


def _extract_excel_sample(path: Path) -> Optional[str]:
    """Extract data from first sheet of Excel file."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        sheet = wb.active
        rows = []
        for i, row in enumerate(sheet.iter_rows(max_row=20, values_only=True)):
            if i >= 20:
                break
            row_str = " | ".join(str(cell) if cell else "" for cell in row[:15])
            rows.append(row_str)
        wb.close()
        return f"Sheet: {sheet.title}\n" + "\n".join(rows)
    except Exception:
        return None


def _extract_csv_sample(path: Path) -> Optional[str]:
    """Extract first 20 rows of CSV."""
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            lines = []
            for i, line in enumerate(f):
                if i >= 20:
                    break
                lines.append(line.strip())
            return "\n".join(lines)
    except Exception:
        return None


def _extract_docx_sample(path: Path) -> Optional[str]:
    """Extract text from Word document."""
    try:
        from docx import Document
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs[:30] if p.text.strip()]
        return "\n".join(paragraphs)
    except Exception:
        return None


def _extract_text_sample(path: Path) -> Optional[str]:
    """Extract from plain text file."""
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read(5000)
    except Exception:
        return None


# =============================================================================
# Utility Functions
# =============================================================================

def get_classification_summary(inventory: List[DocumentInventoryItem]) -> dict:
    """
    Generate classification summary statistics.

    Returns:
        Dict with counts by type, confidence distribution, etc.
    """
    summary = {
        "total": len(inventory),
        "by_type": {},
        "by_source": {},
        "by_confidence": {
            "high": 0,      # >= 0.8
            "medium": 0,    # 0.5 - 0.8
            "low": 0,       # < 0.5
        },
        "unknown_count": 0,
    }

    for item in inventory:
        # By type
        doc_type = item["document_type"]
        summary["by_type"][doc_type] = summary["by_type"].get(doc_type, 0) + 1

        # By source
        source = item["classification_source"]
        summary["by_source"][source] = summary["by_source"].get(source, 0) + 1

        # By confidence
        conf = item["classification_confidence"]
        if conf >= 0.8:
            summary["by_confidence"]["high"] += 1
        elif conf >= 0.5:
            summary["by_confidence"]["medium"] += 1
        else:
            summary["by_confidence"]["low"] += 1

        # Unknown count
        if doc_type == "unknown":
            summary["unknown_count"] += 1

    return summary
